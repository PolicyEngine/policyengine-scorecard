"""Stage HM Treasury's Budget 2026 Tables 6.1, 6.2 and 6.4, cell by cell (#230).

Every value is read from the published XLSX workbooks as stored (openpyxl,
cached values) and staged verbatim: value_raw is the cell, value is the cell
times 1,000,000. Nothing is derived, re-signed or summed. Each workbook is
hashed and the hash must match the manifest pin before a row is written, so a
re-published file stops the stager instead of silently re-staging.

    uv run --with openpyxl python pipeline/stage_ab2026_hmt_tables.py \
        --t61 PATH/Table_6.1_and_6.2_Budget_2026_Policy_Decisions.xlsx \
        --t64 PATH/Table_6.4_Recostings_Table.xlsx

The output gzip is written with mtime 0, so a re-run on the same workbooks
is byte-identical and COUNTS.json's claims_sha256 pin holds.
"""

import argparse
import gzip
import hashlib
import io
import json
import re
import sys
from pathlib import Path

try:
    from pipeline import ab2026_measure_keys as K
except ModuleNotFoundError:  # direct script invocation without PYTHONPATH
    import ab2026_measure_keys as K

ROOT = Path(__file__).resolve().parent.parent
FAMILY = ROOT / "sources/harvest-uk-ab2026-2026-10-29/uk_hmt_tables"
FISCAL_EVENT = "autumn_budget_2026"

T61 = {
    "url": "https://assets.publishing.service.gov.uk/media/6ae187a2779378370fd80a9a/Table_6.1_and_6.2_Budget_2026_Policy_Decisions.xlsx",
    "title": "Table 6.1 and 6.2 Budget 2026 Policy Decisions",
    "sha256": "359958df6b1ab9746073725f6f8017ebdbb55040e57cf28612891f1e294a7f15",
    "bytes": 67126,
}
T64 = {
    "url": "https://assets.publishing.service.gov.uk/media/6ae0f93fe4367e5c8b0467d6/Table_6.4_Recostings_Table.xlsx",
    "title": "Table 6.4 Measures announced at Spring Statement 2026 or earlier that will take effect from 29 October 2026 or later",
    "sha256": "b0449df1737f128fdc32ffeffbe7ed11876e7ab1c80d14dd9908e68a38386d3e",
    "bytes": 76402,
}
PUBLISHED = "2026-10-28"
RETRIEVED = "2026-10-28"
REVERIFIED = "2026-10-29"
COLLECTION = (
    "https://www.gov.uk/government/publications/supporting-documents-for-budget-2026"
)

SIGN = (
    "positive = yield to the Exchequer (reduces borrowing); negative = cost to "
    "the Exchequer (the tables print no legend; convention observed from the "
    "rows themselves, e.g. the bank surcharge rise is + and the Household "
    "Energy Payment is -, as in the AB2025 Table 4.1 harvest)"
)
FT_SIGN = (
    "positive = cash inflow to the public sector; negative = cash outflow "
    "(financial transactions: the table prints no legend; convention observed "
    "from the rows, e.g. the National Wealth Fund's investment is -)"
)


def fy_end(fy):
    m = re.fullmatch(r"(\d{4})-(\d{2})", fy)
    assert m, fy
    return int(m.group(1)) + 1


def raw(v):
    """The cell as stored, without Python's float decoration."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def publication(src, locator):
    return {
        "title": src["title"],
        "url": src["url"],
        "date": PUBLISHED,
        "publisher": "HM Treasury",
        "primary_sha256": src["sha256"],
        "locator": locator,
        "access": "direct",
    }


def base_row(src, *, title, fy, value, locator, table, column, conditions, metric):
    r = {
        "source": "hm_treasury",
        "source_model": "hmt_costing_obr_certified",
        "value": int(value) * 1_000_000 if float(value).is_integer() else value * 1e6,
        "value_raw": raw(value),
        "normalization": "£ million as printed in the cell x 1,000,000; sign as printed",
        "value_kind": "point",
        "period": fy_end(fy),
        "time_basis": "fiscal_year",
        "conditions": {
            "fiscal_event": FISCAL_EVENT,
            "fy": fy,
            "geography": "UK",
            "basis": "forecast",
            **conditions,
        },
        "reform_hint": title,
        "attribution": "own",
        "benchmark_class": "different_model",
        "baseline_policy": None,
        "publication": publication(src, locator),
        "source_table": table,
        "source_column": column,
        "quote": f"{title} | {column}: {raw(value)}",
        "parse_confidence": "high",
        "status": "ok",
        "note": "",
    }
    r.update(metric)
    return r


REVENUE = {"metric": "revenue_change", "unit_concept": "gbp"}
FT = {"proposed_metric": "financial_transaction_cash_impact", "unit_concept": "gbp"}


def header_years(ws, row):
    cols = {}
    for c in ws[row]:
        if isinstance(c.value, str) and re.fullmatch(r"\d{4}-\d{2}", c.value.strip()):
            cols[c.column] = c.value.strip()
    assert len(cols) == 6, (ws.title, cols)
    return cols


def stage_t61(wb, out):
    ws = wb["Table 6.1"]
    years = header_years(ws, 4)
    # Headings sit in column B with nothing beside them. A run of headings
    # with no measure line between them is recorded whole ("A > B"): the
    # table marks no hierarchy beyond that adjacency.
    run, heading = [], ""
    seen = set()
    for row in ws.iter_rows(min_row=5):
        b, c = row[1].value, row[2].value
        if (
            isinstance(b, str)
            and c is None
            and not any(isinstance(x.value, (int, float)) for x in row[3:])
        ):
            text = " ".join(b.split())
            if text[0].isdigit():
                break  # footnotes
            run.append(text)
            heading = " > ".join(run)
            continue
        run = []
        if isinstance(b, int) and isinstance(c, str):
            line, title = b, " ".join(c.split())
            fiscal_head = (row[3].value or "").strip() or None
            head = row[4].value.strip()
            seen.add(line)
            for col, fy in years.items():
                v = ws.cell(row=row[0].row, column=col).value
                assert isinstance(v, (int, float)), (line, fy, v)
                cond = {
                    "sign_convention": SIGN,
                    "scoring_method": "final",
                    "table": "6.1",
                    "table_line": str(line),
                    "head": head,
                    "heading": heading,
                    "row_type": "measure",
                }
                if fiscal_head:
                    cond["fiscal_head"] = fiscal_head
                out.append(
                    {
                        **base_row(
                            T61,
                            title=title,
                            fy=fy,
                            value=v,
                            locator=f"Table 6.1 line {line}, {fy} (cell {ws.cell(row=row[0].row, column=col).coordinate})",
                            table="Table 6.1 Budget 2026 policy decisions (£ million)",
                            column=fy,
                            conditions=cond,
                            metric=REVENUE,
                        ),
                        "measure_key": K.key(line),
                    }
                )
            continue
        if not b and isinstance(c, str) and c.strip().startswith("Total"):
            label = c.strip()
            slug = {
                "Total policy decisions": "package_total_policy_decisions",
                "Total spending policy decisions": "package_total_spending_policy_decisions",
                "Total tax policy decisions": "package_total_tax_policy_decisions",
            }[label]
            for col, fy in years.items():
                v = ws.cell(row=row[0].row, column=col).value
                out.append(
                    {
                        **base_row(
                            T61,
                            title=label,
                            fy=fy,
                            value=v,
                            locator=f"Table 6.1 '{label}', {fy} (cell {ws.cell(row=row[0].row, column=col).coordinate})",
                            table="Table 6.1 Budget 2026 policy decisions (£ million)",
                            column=fy,
                            conditions={
                                "sign_convention": SIGN,
                                "scoring_method": "final",
                                "table": "6.1",
                                "row_type": "total",
                            },
                            metric=REVENUE,
                        ),
                        "measure_key": K.package_key(slug),
                        "note": "Table footnote 1: totals may not sum due to rounding.",
                    }
                )
    assert seen == set(K.TABLE_61), sorted(set(K.TABLE_61) ^ seen)


def stage_t62(wb, out):
    ws = wb["Table 6.2"]
    years = header_years(ws, 4)
    seen = set()
    for row in ws.iter_rows(min_row=5):
        b, c = row[1].value, row[2].value
        if isinstance(b, int) and isinstance(c, str):
            line, title = b, " ".join(c.split())
            seen.add(line)
            fiscal_head = (row[3].value or "").strip()
            for col, fy in years.items():
                v = ws.cell(row=row[0].row, column=col).value
                assert isinstance(v, (int, float)), (line, fy, v)
                cond = {
                    "sign_convention": FT_SIGN,
                    "table": "6.2",
                    "table_line": str(line),
                    "row_type": "measure",
                }
                if fiscal_head:
                    cond["fiscal_head"] = fiscal_head
                out.append(
                    {
                        **base_row(
                            T61,
                            title=title,
                            fy=fy,
                            value=v,
                            locator=f"Table 6.2 line {line}, {fy} (cell {ws.cell(row=row[0].row, column=col).coordinate})",
                            table="Table 6.2 Budget 2026 Financial Transactions (cash impact, £ million)",
                            column=fy,
                            conditions=cond,
                            metric=FT,
                        ),
                        "measure_key": K.PREFIX + K.TABLE_62[line],
                    }
                )
    assert seen == set(K.TABLE_62), seen


def stage_t64(wb, out):
    ws = wb.active
    years = header_years(ws, 3)
    group = None
    seen = set()
    for row in ws.iter_rows(min_row=4):
        f, g = row[5].value, row[6].value
        if isinstance(f, str) and f.startswith("Measures announced at"):
            group = f.strip()
            continue
        if isinstance(f, int) and isinstance(g, str):
            line, title = f, " ".join(g.replace("\xa0", " ").split())
            seen.add(line)
            event, mkey = K.TABLE_64[line]
            head = row[7].value.strip()
            for col, fy in years.items():
                v = ws.cell(row=row[0].row, column=col).value
                assert isinstance(v, (int, float)), (line, fy, v)
                out.append(
                    {
                        **base_row(
                            T64,
                            title=title,
                            fy=fy,
                            value=v,
                            locator=f"Table 6.4 line {line}, {fy} (cell {ws.cell(row=row[0].row, column=col).coordinate})",
                            table="Table 6.4 Measures announced at Spring Statement 2026 or earlier that will take effect from 29 October 2026 or later (£)",
                            column=fy,
                            conditions={
                                "sign_convention": SIGN,
                                "scoring_method": "final",
                                "table": "6.4",
                                "table_line": str(line),
                                "head": head,
                                "row_type": "recosting",
                                "recosting_of": event,
                                "announced_at": group,
                            },
                            metric=REVENUE,
                        ),
                        "measure_key": mkey,
                        "note": (
                            "A recosting at Budget 2026 of a measure announced at an "
                            "earlier event; measure_key is that event's registry key. "
                            "The table's units header prints '(£)', but its "
                            "footnotes and magnitudes are those of Table 6.1 "
                            "(£ million), which this row assumes."
                        ),
                        "parse_confidence": "medium",
                    }
                )
    assert seen == set(K.TABLE_64), seen


def write_gz(rows, path):
    buf = io.BytesIO()
    with gzip.GzipFile(fileobj=buf, mode="wb", mtime=0, filename="") as gz:
        for r in rows:
            gz.write(
                (json.dumps(r, sort_keys=True, ensure_ascii=False) + "\n").encode()
            )
    path.write_bytes(buf.getvalue())


def main(argv=None):
    import openpyxl

    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--t61", required=True)
    ap.add_argument("--t64", required=True)
    args = ap.parse_args(argv)
    for path, src in ((args.t61, T61), (args.t64, T64)):
        got = sha256(path)
        if got != src["sha256"]:
            raise SystemExit(f"{path}: sha256 {got} != pinned {src['sha256']}")
    rows = []
    wb = openpyxl.load_workbook(args.t61, data_only=True)
    stage_t61(wb, rows)
    stage_t62(wb, rows)
    stage_t64(openpyxl.load_workbook(args.t64, data_only=True), rows)
    FAMILY.mkdir(parents=True, exist_ok=True)
    write_gz(rows, FAMILY / "claims_staged.jsonl.gz")
    manifest = [
        {
            "url": src["url"],
            "title": src["title"],
            "date": PUBLISHED,
            "publisher": "HM Treasury",
            "sha256": src["sha256"],
            "bytes": src["bytes"],
            "doc_type": "xlsx",
            "access": "direct",
            "retrieved": RETRIEVED,
            "note": (
                f"Attachment of {COLLECTION}. Downloaded {RETRIEVED} and again "
                f"{REVERIFIED} with a browser User-Agent; both downloads hash to "
                "the pinned sha256."
            ),
        }
        for src in (T61, T64)
    ]
    (FAMILY / "manifest.jsonl").write_text(
        "".join(json.dumps(m, ensure_ascii=False) + "\n" for m in manifest)
    )
    by = {}
    for r in rows:
        t = r["conditions"]["table"]
        by[t] = by.get(t, 0) + 1
    print(f"staged {len(rows)} rows: {by}")


if __name__ == "__main__":
    main(sys.argv[1:])
