"""Read-only source/workbook receipts for POPULATIONS_AND_VINTAGES.md.

No engine imports, simulations, dataset downloads or builds. Run in the existing
research environment while /tmp/uk-replay-scope is available. Public GOV.UK
metadata requests are optional (--refresh-govuk); these fetch JSON, not datasets.
"""

import argparse
import collections
import datetime
import hashlib
import json
from pathlib import Path
import re
import subprocess
import time
import urllib.request

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
SCRATCH = Path("/tmp/uk-replay-scope")
HARVEST = Path.home() / "scorecard-harvest"
OUT = Path(__file__).with_name("population_receipts.json")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize(value):
    if isinstance(value, datetime.datetime):
        return value.strftime("%B %Y")
    match = re.match(r"^(\w+)\s*-?\s*(20\d{2})", str(value))
    return " ".join(match.groups()) if match else str(value)


def collect():
    receipt = {"checked_on": "2026-10-09", "method": "source/workbook inspection only", "code": []}
    files = {
        "microcosm": {
            "packages/microcosm-build/src/microcosm/build/uk/frs_release.json": [(1, 26)],
            "packages/microcosm-build/src/microcosm/build/uk/chronicle_feed.json": [(1, 22)],
            "packages/microcosm-build/src/microcosm/build/uk_runtime/frs_release.py": [(45, 52), (108, 121)],
            "packages/microcosm-build/src/microcosm/build/uk_runtime/country_adapter.py": [(35, 90)],
            "packages/microcosm-build/src/microcosm/build/uk_runtime/full_targets.py": [(127, 153)],
            "packages/microcosm-build/src/microcosm/build/uk_runtime/graph_calibration.py": [(1, 39), (112, 164)],
            "packages/microcosm-build/src/microcosm/build/ledger_targets.py": [(35, 40), (3180, 3240), (3270, 3290), (3717, 3770)],
            "packages/microcosm-calibrate/src/microcosm/calibrate/static_aging.py": [(336, 353)],
            "packages/microcosm-data/src/microcosm/data/registry.py": [(101, 154)],
            "packages/microcosm-data/src/microcosm/data/annual_projections.py": [(1, 12)],
            "docs/static-aging.md": [(1, 37), (113, 142)],
            "docs/uk-full-build-graph.md": [(1, 59), (88, 106)],
            "changelog.d/723-uk-frs-2024-25-retarget.changed.md": [(1, 4)],
        },
        "chronicle": {
            "chronicle/consumer_contract.py": [(90, 109)],
            "docs/schemas/consumer_fact.v4.schema.json": [(260, 292), (389, 435)],
            "packages/obr/efo_receipts_march_2026/source_package.yaml": [(1, 29), (70, 88)],
            "README.md": [(24, 58)],
        },
        "policyengine-uk-data": {
            "policyengine_uk_data/targets/sources.yaml": [(1, 40)],
            "policyengine_uk_data/targets/sources/obr.py": [(1, 39), (45, 61)],
        },
    }
    for repo, paths in files.items():
        commit = subprocess.check_output(["git", "-C", str(SCRATCH / repo), "rev-parse", "HEAD"], text=True).strip()
        for path, spans in paths.items():
            file = SCRATCH / repo / path
            lines = file.read_text().splitlines()
            receipt["code"].append({"repo": "PolicyEngine/" + repo, "commit": commit, "path": path,
                                    "sha256": digest(file), "snippets": [{"first_line": a, "last_line": b,
                                    "text": "\n".join(lines[a-1:b])} for a, b in spans]})
    uk = SCRATCH / "microcosm/packages/microcosm-build/src/microcosm/build/uk"
    refs = json.loads((uk / "target_references.json").read_text())["target_references"]
    parity = json.loads((uk / "ledger_compile_parity_production_2023_signed_differences.json").read_text())
    receipt["microcosm_target_surface"] = {"reference_count": len(refs), "compiled_2023": parity["compiled_count"],
                                           "parity_file_sha256": digest(uk / "ledger_compile_parity_production_2023_signed_differences.json")}
    fixture = SCRATCH / "microcosm/packages/microcosm-data/tests/fixtures/uk_june_2023/populace-uk-2023-dd68c73-4aa4b14-20260619T023711Z"
    release = json.loads((fixture / "release_manifest.json").read_text())
    build = json.loads((fixture / "build_manifest.json").read_text())
    diagnostics = json.loads((fixture / "calibration_diagnostics.json").read_text())
    receipt["certified_2023_fixture"] = {"relative_path": str(fixture.relative_to(SCRATCH / "microcosm")),
                                         "build": release["build"], "dataset": build["dataset"], "code": build["code"],
                                         "calibration": build["calibration"], "options": diagnostics["options"],
                                         "files_sha256": {name: digest(fixture / name) for name in
                                         ["release_manifest.json", "build_manifest.json", "calibration_diagnostics.json"]}}
    obr_packages = SCRATCH / "chronicle/packages/obr"
    receipt["chronicle_obr_packages"] = [{"path": str(file.relative_to(SCRATCH / "chronicle")), "sha256": digest(file),
                                        "vintages": sorted(set(re.findall(r"^\s+vintage: (.+)$", file.read_text(), re.M)))}
                                        for file in sorted(obr_packages.glob("*/source_package.yaml"))]

    rows = [json.loads(line) for line in (ROOT / "data/uk/events/base_rows.jsonl").read_text().splitlines()]
    events = sorted({(row["event_index"], row["event"]) for row in rows})
    vintages = ["June 2010", "November 2010", "March 2011", "November 2011", "March 2012", "December 2012",
                "March 2013", "December 2013", "March 2014", "December 2014", "March 2015", "July 2015",
                "November 2015", "March 2016", "November 2016", "March 2017", "November 2017", "March 2018",
                "October 2018", "March 2019", "March 2020", "November 2020", "March 2021", "October 2021",
                "March 2022", "November 2022", "March 2023", "November 2023", "March 2024", "October 2024",
                "March 2025", "November 2025"]
    assert len(events) == len(vintages) == 32
    receipt["event_vintages"] = [{"index": index, "event": event, "vintage": vintage,
                                   "costing_years": sorted(set().union(*(set(row["costing_gbp_m"]) for row in rows if row["event"] == event)))}
                                   for (index, event), vintage in zip(events, vintages)]
    receipt["hofd"] = {}
    for edition in ["March_2025", "Spring_2026"]:
        file = HARVEST / "uk_obr/downloads" / f"Historical_official_forecasts_database_{edition}.xlsx"
        wb = openpyxl.load_workbook(file, data_only=True, read_only=True)
        sheets = [sheet for sheet in wb if sheet.title not in ["Index", "Contents", "Aggregates", "Receipts", "Spending", "Economy"]
                  and not sheet.title.endswith("(2)")]
        extracted = {}
        for ws in sheets:
            # Numeric values are tied to the workbook's year header, not cell position.
            headers = {cell.column: str(cell.value) for cell in ws[4]
                       if re.match(r"^(19|20)\d{2}(-\d{2})?$", str(cell.value))}
            parsed = []
            for row in ws.iter_rows(min_row=5, max_row=200, max_col=65):
                label = normalize(row[0].value)
                if label not in vintages and not str(row[0].value).startswith("Outturn data"):
                    continue
                values = {headers[cell.column]: {"cell": cell.coordinate, "value": cell.value}
                          for cell in row if getattr(cell, "column", None) in headers and isinstance(cell.value, (int, float))}
                if values:
                    parsed.append({"row": row[0].row, "vintage": label, "raw_label": str(row[0].value), "values": values})
            extracted[ws.title] = parsed
        # Preserve only replay target/economic rows; compare all historic cells.
        selected = ["IT", "SA IT", "PAYE IT", "NICS", "HSC", "VAT", "Fuel", "CGT", "IHT", "PTT", "Council", "Total welfare",
                    "Welfare in", "Welfare out", "CPI", "RPI", "Earnings", "Wages&Salaries", "Empl", "Unemplrate", "Houseprices", "CC"]
        available = [name for name in selected if name in extracted]
        receipt["hofd"][edition] = {"file": str(file), "sha256": digest(file), "sheet_count": len(wb.sheetnames),
                                    "data_sheet_count": len(sheets), "contents_B3": wb["Contents"]["B3"].value,
                                    "target_sheets": {name: extracted[name] for name in available},
                                    "all_sheets": list(extracted), "all_cells": extracted}
    old = receipt["hofd"]["March_2025"]["all_cells"]
    new = receipt["hofd"]["Spring_2026"]["all_cells"]
    differences, common = [], 0
    for name, oldrows in old.items():
        newrows = {row["vintage"]: row for row in new[name]}
        for row in oldrows:
            if row["vintage"].startswith("Outturn") or row["vintage"] not in newrows:
                continue
            other = newrows[row["vintage"]]
            for year, cell in row["values"].items():
                if year not in other["values"]:
                    differences.append({"sheet": name, "vintage": row["vintage"], "year": year,
                                        "old": cell["value"], "new": None})
                    continue
                common += 1
                value = other["values"][year]["value"]
                if cell["value"] != value:
                    differences.append({"sheet": name, "vintage": row["vintage"], "year": year,
                                        "old": cell["value"], "new": value})
    receipt["hofd_comparison"] = {"shared_numeric_cells": common, "differences_including_removed_cells": differences}
    for edition in receipt["hofd"].values():
        del edition["all_cells"]
    receipt["detailed_workbooks"] = []
    prior_detailed = {}
    if OUT.exists():
        prior_detailed = {item["file"]: item for item in json.loads(OUT.read_text()).get("detailed_workbooks", [])}
    for file in sorted((HARVEST / "uk_obr/downloads").glob("*.xlsx")):
        if not any(prefix in file.name.lower() for prefix in ["receipts", "expenditure", "ready_reckoner"]):
            continue
        if str(file) in prior_detailed and digest(file) == prior_detailed[str(file)]["sha256"]:
            receipt["detailed_workbooks"].append(prior_detailed[str(file)])
            continue
        wb = openpyxl.load_workbook(file, data_only=True, read_only=True)
        evidence = {"file": str(file), "sha256": digest(file), "sheets": wb.sheetnames, "samples": []}
        for ws in wb:
            title = str(ws["B2"].value)
            rows_to_keep = []
            for row_number, row in enumerate(ws.iter_rows(max_row=300, max_col=50), start=1):
                text = " ".join(str(cell.value) for cell in row[:3] if cell.value is not None)
                if re.search(r"Class 1|Class 2|Class 4|Population \(|Universal credit|Universal Credit|State pension|State Pension|post-measures|earnings|March 2025 EFO", text):
                    rows_to_keep.append({"row": row_number, "cells": {cell.coordinate: cell.value for cell in row if cell.value is not None}})
            if rows_to_keep:
                evidence["samples"].append({"sheet": ws.title, "title_B2": title, "rows": rows_to_keep[:15]})
        receipt["detailed_workbooks"].append(evidence)
    # Preserve public archive file identities already researched; their status is
    # separately labelled as prior-lane metadata, with new spot checks in the doc.
    archive = SCRATCH / "agents/work_archive_obr/curated.json"
    receipt["obr_archive_catalog"] = {"basis": "prior-lane CDX/file index, reviewed; not all bytes re-fetched",
                                      "catalog_sha256": digest(archive), "events": json.loads(archive.read_text())}
    return receipt


def refresh_govuk(receipt):
    pages = ["benefit-expenditure-and-caseload-tables-2010",
             "data-about-people-that-were-receiving-benefits-in-2011-as-published-with-the-budget-and-the-autumn-statement"]
    pages += [f"benefit-expenditure-and-caseload-tables-{year}" for year in range(2012, 2026)]
    receipt["dwp_public_metadata"] = []
    for slug in pages:
        url = "https://www.gov.uk/api/content/government/publications/" + slug
        try:
            request = urllib.request.Request(url, headers={"User-Agent": "policyengine-scorecard-research/1.0"})
            with urllib.request.urlopen(request, timeout=25) as response:
                raw = response.read()
            data = json.loads(raw)
            detail = data.get("details", {})
            receipt["dwp_public_metadata"].append({"page": "https://www.gov.uk/government/publications/" + slug,
                 "api_sha256": hashlib.sha256(raw).hexdigest(), "first_published_at": data.get("first_published_at"),
                 "attachments": [{key: att.get(key) for key in ["title", "url", "content_type"]}
                 for att in detail.get("attachments", []) if re.search("outturn.*forecast", att.get("title", ""), re.I)],
                 "change_history": detail.get("change_history", [])})
        except Exception as exc:
            receipt["dwp_public_metadata"].append({"page": url, "error": str(exc), "availability": "unknown"})
        time.sleep(1)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--refresh-govuk", action="store_true")
    args = parser.parse_args()
    receipt = collect()
    if args.refresh_govuk:
        refresh_govuk(receipt)
    elif OUT.exists():
        prior = json.loads(OUT.read_text())
        if "dwp_public_metadata" in prior:
            receipt["dwp_public_metadata"] = prior["dwp_public_metadata"]
    OUT.write_text(json.dumps(receipt, indent=2, default=str) + "\n")
    print(json.dumps({"output": str(OUT), "bytes": OUT.stat().st_size,
                      "events": len(receipt["event_vintages"]), "hofd_comparison": receipt["hofd_comparison"]}))
