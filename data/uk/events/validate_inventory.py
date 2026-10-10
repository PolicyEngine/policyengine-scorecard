"""Engine-free inventory accounting and optional source-workbook verification.

Run with the existing environment (openpyxl needed only for --workbook):
  python data/uk/events/validate_inventory.py
  python data/uk/events/validate_inventory.py --workbook PATH --receipt PATH
No simulations or population data are loaded.
"""
import argparse
from collections import Counter, defaultdict
import hashlib
import json
from pathlib import Path

from build_inventory import HERE, EVENTS, HARVEST, load_rows, build


def check(workbook=None):
    rows = [json.loads(line) for line in (HERE / 'inventory.jsonl').read_text().splitlines()]
    assert len(rows) == 3796 and len({r['row_id'] for r in rows}) == len(rows)
    assert {r['event'] for r in rows} == set(EVENTS)
    expected, summary = build(False)
    assert (HERE / 'inventory.jsonl').read_text() == expected
    assert (HERE / 'event_summary.json').read_text() == summary
    source = load_rows()
    assert sum(len(r['costing_gbp_m']) for r in source) == sum(len(r['costing_gbp_m']) for r in rows)
    for raw, row in zip(source, rows):
        assert raw['costing_gbp_m'] == row['costing_gbp_m']
        assert row['scope_note']
        if row['scope'] == 'in': assert row['expressibility_note']
        assert row['scope'] == ('out' if row['class'] == 'out' else 'in')
        assert row['readiness'] == 'scoped_only_no_vintage_replay_certified'
        if row['scope'] == 'in': assert row['mechanism_evidence']
    by_key = Counter((r['event'], r['table'], r['measure'], r['head']) for r in rows)
    duplicate_occurrences = sum(n - 1 for n in by_key.values())
    receipt = {'rows': len(rows), 'events': len(EVENTS),
               'source_cells': sum(len(r['costing_gbp_m']) for r in rows),
               'zero_cells': sum(v == 0 for r in rows for v in r['costing_gbp_m'].values()),
               'repeated_key_occurrences_preserved': duplicate_occurrences,
               'inventory_sha256': hashlib.sha256(expected.encode()).hexdigest(),
               'checks': ['deterministic rebuild', 'unique occurrence ids', 'all 32 events',
                          'source FY values preserved', 'scope and mechanism metadata']}
    if not workbook: return receipt
    import openpyxl
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    wanted = set(EVENTS); actual = []; locations = []; extensions = 0
    for sheet, table in [('Tax Measures', 'tax'), ('Spending Measures', 'spending')]:
        ws = book[sheet]
        headers = next(ws.iter_rows(min_row=3, max_row=3))
        fy_columns = [(i, c.value) for i, c in enumerate(headers)
                      if isinstance(c.value, str) and len(c.value) == 7 and c.value[4] == '-']
        for number, cells in enumerate(ws.iter_rows(min_row=4), 4):
            if cells[1].value not in wanted or not cells[2].value: continue
            values = {}; columns = {}
            for col, fy in fy_columns:
                c = cells[col]
                if c.value is None or c.value == '': continue
                try: value = float(c.value)
                except (TypeError, ValueError): continue
                if c.fill.patternType == 'solid' and c.fill.start_color.rgb == 'FFE1E9EE':
                    extensions += 1; continue
                values[fy] = value; columns[fy] = c.column_letter
            if values:
                actual.append((cells[1].value, table, str(cells[2].value),
                               str(cells[3].value) if cells[3].value else '', values))
                locations.append({'sheet': sheet, 'row': number, 'fy_columns': columns})
    expected_rows = [(r['event'], r['table'], r['measure'], r['head'], r['costing_gbp_m']) for r in rows]
    assert len(actual) == len(expected_rows), (len(actual), len(expected_rows))
    for i, (a, e) in enumerate(zip(actual, expected_rows)):
        assert a == e, (i, a, e)
        locations[i]['row_id'] = rows[i]['row_id']
    receipt.update({'workbook': Path(workbook).name,
                    'workbook_sha256': hashlib.sha256(Path(workbook).read_bytes()).hexdigest(),
                    'workbook_all_original_cells_match': True,
                    'excluded_extended_cells': extensions, 'workbook_locations': locations})
    book.close()
    return receipt


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--workbook', type=Path)
    parser.add_argument('--receipt', type=Path)
    args = parser.parse_args()
    result = check(args.workbook)
    if args.receipt: args.receipt.write_text(json.dumps(result, indent=1, sort_keys=True) + '\n')
    print(json.dumps({k:v for k,v in result.items() if k != 'workbook_locations'}, indent=1))
